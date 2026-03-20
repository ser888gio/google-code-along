"""
Generates a Discovery Session Word document (.docx) from structured data.

Takes the JSON produced by compile_srs_json_from_session and writes
each section into a properly formatted Word document.
"""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import RGBColor


def _set_heading_color(paragraph, color: tuple[int, int, int]):
    """Set the font color on all runs in a heading paragraph."""
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(*color)


def _add_section_heading(doc: Document, text: str, level: int):
    """Add a numbered heading at the given level (1-4)."""
    p = doc.add_heading(text, level=level)
    colors = {
        1: (26, 58, 42),  # dark green
        2: (42, 74, 58),
        3: (60, 100, 80),
        4: (80, 120, 100),
    }
    _set_heading_color(p, colors.get(level, (0, 0, 0)))
    return p


def _add_placeholder(doc: Document, text: str):
    """Add body text. If text is empty/None, add a TBD note."""
    content = (text or "").strip()
    if not content:
        p = doc.add_paragraph()
        run = p.add_run("TBD — information to be gathered.")
        run.italic = True
        run.font.color.rgb = RGBColor(150, 150, 150)
    else:
        doc.add_paragraph(content)


def srs_json_to_docx(srs: dict) -> bytes:
    """Convert a structured Discovery Session dict to a .docx file and return the bytes.

    Expected structure has four top-level blocks:
      context_and_business_problem, target_state_and_success,
      scope_and_boundaries, data_systems_integrations_access
    Each value is a dict with sub-fields.
    """
    doc = Document()

    # ── Title ──────────────────────────────────────────────────────────────────
    title_para = doc.add_heading("Discovery Session — Requirement Specification", 0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_para.runs:
        run.font.color.rgb = RGBColor(26, 58, 42)

    doc.add_paragraph()

    # ── 1. Context & Business Problem ──────────────────────────────────────────
    block1 = srs.get("context_and_business_problem", {})
    _add_section_heading(doc, "1. Context & Business Problem", 1)

    _add_section_heading(doc, "1.1 Context", 2)
    _add_placeholder(doc, block1.get("context"))

    _add_section_heading(doc, "1.2 Business Problem", 2)
    _add_placeholder(doc, block1.get("business_problem"))

    # ── 2. Target State & Success ──────────────────────────────────────────────
    block2 = srs.get("target_state_and_success", {})
    _add_section_heading(doc, "2. Target State & Success", 1)

    _add_section_heading(doc, "2.1 Target State", 2)
    _add_placeholder(doc, block2.get("target_state"))

    _add_section_heading(doc, "2.2 Success Criteria", 2)
    _add_placeholder(doc, block2.get("success"))

    # ── 3. Scope & Boundaries ─────────────────────────────────────────────────
    block3 = srs.get("scope_and_boundaries", {})
    _add_section_heading(doc, "3. Scope & Boundaries", 1)

    _add_section_heading(doc, "3.1 Scope — In/Out", 2)
    _add_placeholder(doc, block3.get("scope_in_out"))

    _add_section_heading(doc, "3.2 Boundaries & Dependencies", 2)
    _add_placeholder(doc, block3.get("boundaries"))

    # ── 4. Data, Systems, Integrations & Access ───────────────────────────────
    block4 = srs.get("data_systems_integrations_access", {})
    _add_section_heading(doc, "4. Data, Systems, Integrations & Access", 1)

    _add_section_heading(doc, "4.1 Data & Systems", 2)
    _add_placeholder(doc, block4.get("data_systems"))

    _add_section_heading(doc, "4.2 Integrations & Access", 2)
    _add_placeholder(doc, block4.get("integrations_access"))

    # ── Serialize ─────────────────────────────────────────────────────────────
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
